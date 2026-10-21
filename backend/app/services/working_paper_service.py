import io
import os
import re
from datetime import date, datetime

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from pymongo import DESCENDING, ReturnDocument

from app.db.session import db
from app.services.cycle_service import get_mission_cycle_detail
from app.services.document_service import store_upload
from app.services.mission_service import get_mission_or_404
from app.services.revision_referentiel import ETAPES
from app.services.revision_service import (
    auteur,
    get_revision_doc,
    now_iso,
    verifier_code,
    verifier_edition,
)

VERSIONS_CONSERVEES = 50
EXTRAIT_MAX_LIGNES = 200
EXTRAIT_MAX_COLONNES = 30


def _feuilles_lead(mission_id, code):
    colonnes = ["Compte", "Libellé", "Solde N", "Solde N-1", "Variation", "Commentaire"]
    try:
        comptes = get_mission_cycle_detail(mission_id, code, None, None)["comptes"]
    except HTTPException:
        # Balances N / N-1 absentes : le PT démarre vide, le collaborateur le complète.
        comptes = []
    data = [
        [c["numero"], c["libelle"], c["soldeN"], c["soldeNMoins1"], f"=C{i}-D{i}", ""]
        for i, c in enumerate(comptes, start=1)
    ]
    if data:
        n = len(data)
        data.append(["Total", "", f"=SUM(C1:C{n})", f"=SUM(D1:D{n})", f"=C{n + 1}-D{n + 1}", ""])
    return [{"nom": "Lead", "colonnes": colonnes, "data": data or [["", "", "", "", "", ""]]}]


def _feuilles_controles():
    return [
        {
            "nom": "Balance âgée",
            "colonnes": ["Tranche", "Montant", "% du solde"],
            "data": [
                ["0-30 jours", "", "=IF(B5=0,0,B1/B5)"],
                ["31-60 jours", "", "=IF(B5=0,0,B2/B5)"],
                ["61-90 jours", "", "=IF(B5=0,0,B3/B5)"],
                ["> 90 jours", "", "=IF(B5=0,0,B4/B5)"],
                ["Total", "=SUM(B1:B4)", ""],
            ],
        },
        {
            "nom": "Cut-off",
            "colonnes": ["N° facture", "Date facture", "Date livraison", "Montant", "Rattachement"],
            "data": [["", "", "", "", ""] for _ in range(10)],
        },
        {
            "nom": "Circularisation",
            "colonnes": ["Client", "Solde comptable", "Montant confirmé", "Écart", "Statut"],
            "data": [["", "", "", f"=B{i}-C{i}", ""] for i in range(1, 11)],
        },
    ]


def _papier_par_defaut(mission_id, code, etape):
    if etape == "4b":
        titre, feuilles = f"PT — Lead du cycle {code}", _feuilles_lead(mission_id, code)
    elif etape == "4d":
        titre, feuilles = f"PT — Contrôles détaillés {code}", _feuilles_controles()
    else:
        titre = f"PT — Cycle {code} · {etape}"
        feuilles = [{"nom": "Feuille 1", "colonnes": ["A", "B", "C", "D", "E"], "data": [["", "", "", "", ""]]}]
    return {
        "titre": titre,
        "mode": "integre",
        "lien": {"provider": "onedrive", "webUrl": "", "chemin": "", "fichier": None, "feuilles": [], "derniereSync": None},
        "feuilles": feuilles,
        "version": 0,
        "createdAt": now_iso(),
        "updatedAt": None,
        "updatedBy": None,
    }


def _verifier_etape(etape):
    if etape not in ETAPES:
        raise HTTPException(status_code=404, detail="Étape inconnue")


def _get_or_create(mission_id, code, etape):
    verifier_code(code)
    _verifier_etape(etape)
    filtre = {"missionId": mission_id, "cycleCode": code, "etape": etape}
    papier = db.working_papers.find_one(filtre)
    if papier is None:
        db.working_papers.update_one(
            filtre, {"$setOnInsert": _papier_par_defaut(mission_id, code, etape)}, upsert=True
        )
        papier = db.working_papers.find_one(filtre)
    return papier


def serialize_papier(papier, revision_doc):
    return {
        "id": str(papier["_id"]),
        "missionId": papier["missionId"],
        "cycleCode": papier["cycleCode"],
        "etape": papier["etape"],
        "titre": papier["titre"],
        "mode": papier["mode"],
        "lien": papier.get("lien") or {},
        "feuilles": papier.get("feuilles") or [],
        "version": papier["version"],
        "updatedAt": papier.get("updatedAt"),
        "updatedBy": papier.get("updatedBy"),
        "verrouille": revision_doc["statut"] not in ETAPES,
    }


def get_papier(mission_id: str, code: str, etape: str):
    get_mission_or_404(mission_id)
    papier = _get_or_create(mission_id, code, etape)
    return serialize_papier(papier, get_revision_doc(mission_id, code))


def _archiver_version(papier):
    db.working_paper_versions.insert_one(
        {
            "paperId": papier["_id"],
            "version": papier["version"],
            "titre": papier["titre"],
            "mode": papier["mode"],
            "lien": papier.get("lien"),
            "feuilles": papier.get("feuilles"),
            "updatedAt": papier.get("updatedAt") or papier.get("createdAt"),
            "updatedBy": papier.get("updatedBy"),
        }
    )
    anciennes = (
        db.working_paper_versions.find({"paperId": papier["_id"]}, {"_id": 1})
        .sort("version", DESCENDING)
        .skip(VERSIONS_CONSERVEES)
    )
    ids = [v["_id"] for v in anciennes]
    if ids:
        db.working_paper_versions.delete_many({"_id": {"$in": ids}})


def _enregistrer(papier, version, sets, employee):
    """Écriture versionnée : refusée si le PT a été modifié depuis la version lue par le client."""
    par, le = auteur(employee), now_iso()
    updated = db.working_papers.find_one_and_update(
        {"_id": papier["_id"], "version": version},
        {"$set": {**sets, "updatedAt": le, "updatedBy": par}, "$inc": {"version": 1}},
        return_document=ReturnDocument.AFTER,
    )
    if updated is None:
        raise HTTPException(
            status_code=409,
            detail="Le papier de travail a été modifié entre-temps par un autre utilisateur. Rechargez-le.",
        )
    _archiver_version(papier)
    return updated


def update_papier(mission_id: str, code: str, etape: str, payload, employee):
    _, mission = get_mission_or_404(mission_id)
    revision = get_revision_doc(mission_id, code)
    verifier_edition(revision, mission, employee)
    papier = _get_or_create(mission_id, code, etape)

    sets = {}
    if payload.titre is not None:
        sets["titre"] = payload.titre.strip() or papier["titre"]
    if payload.mode is not None:
        sets["mode"] = payload.mode
    if payload.lien is not None:
        for champ, valeur in payload.lien.model_dump().items():
            sets[f"lien.{champ}"] = valeur
    if payload.feuilles is not None:
        if not payload.feuilles:
            raise HTTPException(status_code=400, detail="Le classeur doit contenir au moins un onglet.")
        sets["feuilles"] = [f.model_dump() for f in payload.feuilles]

    updated = _enregistrer(papier, payload.version, sets, employee) if sets else papier
    return serialize_papier(updated, revision)


def list_versions(mission_id: str, code: str, etape: str):
    get_mission_or_404(mission_id)
    papier = _get_or_create(mission_id, code, etape)
    versions = db.working_paper_versions.find(
        {"paperId": papier["_id"]}, {"version": 1, "mode": 1, "updatedAt": 1, "updatedBy": 1}
    ).sort("version", DESCENDING)
    journal = [
        {"version": papier["version"], "mode": papier["mode"], "updatedAt": papier.get("updatedAt"), "updatedBy": papier.get("updatedBy"), "courante": True}
    ]
    journal += [
        {"version": v["version"], "mode": v["mode"], "updatedAt": v.get("updatedAt"), "updatedBy": v.get("updatedBy"), "courante": False}
        for v in versions
    ]
    return journal


def _valeur_excel(valeur):
    if isinstance(valeur, str):
        texte = valeur.strip()
        if re.fullmatch(r"-?\d+([.,]\d+)?", texte.replace(" ", "").replace(" ", "")):
            return float(texte.replace(" ", "").replace(" ", "").replace(",", "."))
    return valeur


def _construire_classeur(papier):
    wb = Workbook()
    wb.remove(wb.active)
    for feuille in papier.get("feuilles") or []:
        ws = wb.create_sheet(title=re.sub(r"[\[\]:*?/\\]", "-", feuille["nom"])[:31] or "Feuille")
        for col, titre in enumerate(feuille.get("colonnes") or [], start=1):
            ws.cell(row=1, column=col, value=titre)
        # Dans l'éditeur intégré, la ligne 1 est la première ligne de données (les en-têtes
        # sont des titres de colonnes). Dans le fichier Excel, les en-têtes occupent la
        # ligne 1 : chaque formule est donc décalée d'une ligne vers le bas.
        for r, ligne in enumerate(feuille.get("data") or [], start=1):
            for c, valeur in enumerate(ligne, start=1):
                if isinstance(valeur, str) and valeur.startswith("="):
                    origine = f"{get_column_letter(c)}{r}"
                    cible = f"{get_column_letter(c)}{r + 1}"
                    valeur = Translator(valeur, origin=origine).translate_formula(cible)
                ws.cell(row=r + 1, column=c, value=_valeur_excel(valeur))
    if not wb.sheetnames:
        wb.create_sheet("Feuille 1")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def exporter_papier(mission_id: str, code: str, etape: str):
    """Exporte le PT intégré en .xlsx dans les uploads de la mission, pour le référencer comme
    pièce probante (4c) ou le télécharger."""
    get_mission_or_404(mission_id)
    papier = _get_or_create(mission_id, code, etape)
    nom = re.sub(r"\s+", "_", re.sub(r"[^\w\- ]", "", papier["titre"]).strip()) or f"PT_{code}_{etape}"
    file_name, file_url = store_upload(mission_id, f"{nom}.xlsx", _construire_classeur(papier))
    return {"fileName": file_name, "fileUrl": file_url, "version": papier["version"], "exporteLe": now_iso()}


def _cellule_lisible(valeur):
    if isinstance(valeur, (datetime, date)):
        return valeur.strftime("%d/%m/%Y")
    return valeur


def _extraire_feuilles(content: bytes):
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Fichier Excel illisible (formats acceptés : .xlsx, .xlsm).")
    feuilles = []
    for ws in wb.worksheets:
        lignes = []
        for ligne in ws.iter_rows(max_row=EXTRAIT_MAX_LIGNES + 1, max_col=EXTRAIT_MAX_COLONNES, values_only=True):
            lignes.append([_cellule_lisible(v) for v in ligne])
        while lignes and all(v is None for v in lignes[-1]):
            lignes.pop()
        largeur = max((max((i + 1 for i, v in enumerate(l) if v is not None), default=0) for l in lignes), default=0)
        lignes = [l[:largeur] for l in lignes]
        colonnes = [str(v) if v is not None else "" for v in (lignes[0] if lignes else [])]
        feuilles.append({"nom": ws.title, "colonnes": colonnes, "data": lignes[1:]})
    wb.close()
    return feuilles


def importer_fichier_lie(mission_id: str, code: str, etape: str, version: int, fichier: UploadFile, employee):
    """Mode « Lien Excel » : dépôt de la copie du classeur, dont les valeurs sont relues
    (lecture seule) pour alimenter l'écran. Redéposer le fichier = resynchroniser."""
    _, mission = get_mission_or_404(mission_id)
    revision = get_revision_doc(mission_id, code)
    verifier_edition(revision, mission, employee)
    papier = _get_or_create(mission_id, code, etape)

    if not fichier.filename or os.path.splitext(fichier.filename)[1].lower() not in {".xlsx", ".xlsm"}:
        raise HTTPException(status_code=400, detail="Formats acceptés : .xlsx, .xlsm.")
    content = fichier.file.read()
    feuilles = _extraire_feuilles(content)
    file_name, file_url = store_upload(mission_id, fichier.filename, content)

    updated = _enregistrer(
        papier,
        version,
        {
            "mode": "lien",
            "lien.fichier": {"fileName": file_name, "fileUrl": file_url},
            "lien.feuilles": feuilles,
            "lien.derniereSync": now_iso(),
        },
        employee,
    )
    return serialize_papier(updated, revision)
