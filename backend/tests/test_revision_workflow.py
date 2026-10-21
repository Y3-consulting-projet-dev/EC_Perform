import io

from openpyxl import load_workbook

from app.db.session import db
from app.services.document_service import upload_file_path
from app.services.revision_referentiel import CHECKLISTS


def _url(mission, suffix=""):
    return f"/missions/{mission}/cycles/C{suffix}"


def _patch_section(api, mission, headers, section, data):
    version = api.get(_url(mission, "/revision"), headers=headers).json()["version"]
    return api.patch(
        _url(mission, "/revision"),
        headers=headers,
        json={"section": section, "data": data, "version": version},
    )


def _preparer_et_soumettre(api, mission, headers):
    assert _patch_section(
        api, mission, headers, "priseConnaissance",
        {"processus": "Facturation ERP", "risques": ["Cut-off", "Cut-off", " "], "referentiel": "SYSCOHADA"},
    ).status_code == 200
    assert _patch_section(
        api, mission, headers, "conclusion",
        {"synthese": "Solde cohérent", "ajustements": [], "decision": "justifie"},
    ).status_code == 200
    response = api.post(_url(mission, "/revision/soumettre"), headers=headers)
    assert response.status_code == 200, response.text
    return response.json()


def _cocher_checklist(api, mission, niveau, headers):
    items = {cle: True for cle, _ in CHECKLISTS[niveau]}
    response = api.patch(_url(mission, f"/revue/{niveau}/checklist"), headers=headers, json={"items": items})
    assert response.status_code == 200, response.text


def _etats(revision):
    return {e["niveau"]: e["etat"] for e in revision["chaine"]}


def test_sections_horodatage_et_verrou_optimiste(api, employees, mission):
    assistant = employees["assistant"]["headers"]

    revision = api.get(_url(mission, "/revision"), headers=assistant).json()
    assert revision["statut"] == "4a" and revision["version"] == 0

    response = _patch_section(
        api, mission, assistant, "justification",
        {"assertions": {"existence": True}, "pieces": [{"libelle": "Balance âgée", "statut": "Reçue"}], "commentaire": "RAS"},
    )
    revision = response.json()
    existence = next(a for a in revision["justification"]["assertions"] if a["cle"] == "existence")
    assert existence["coche"] and existence["par"]["nom"] == "Assistant ASSISTANT" and existence["le"]
    assert revision["justification"]["pieces"][0]["id"]
    assert revision["justification"]["commentaire"]["texte"] == "RAS"
    assert revision["statut"] == "4c"

    # Version périmée -> conflit.
    stale = api.patch(
        _url(mission, "/revision"),
        headers=assistant,
        json={"section": "lead", "data": {"commentaires": {}}, "version": 0},
    )
    assert stale.status_code == 409

    # La mission passe en Phase 4 dès qu'un cycle est en révision.
    assert db.missions.find_one()["phase"].startswith("4/7")


def test_seul_le_collaborateur_affecte_ou_l_encadrement_edite(api, employees, mission):
    autre = db.employees.insert_one({"nom": "X", "prenoms": "Y", "grade": "Assistant", "email": "x@test.ci", "password_hash": "x"})
    from app.core.security import create_access_token

    headers = {"Authorization": f"Bearer {create_access_token('x@test.ci')}"}
    assert autre.inserted_id
    response = _patch_section(api, mission, headers, "lead", {"commentaires": {"411000": "ok"}})
    assert response.status_code == 403
    response = _patch_section(api, mission, employees["senior"]["headers"], "lead", {"commentaires": {"411000": "ok"}})
    assert response.status_code == 200


def test_completer_la_justification_ramene_en_4c(api, employees, mission):
    assistant = employees["assistant"]["headers"]
    _patch_section(api, mission, assistant, "constats", {"constats": [{"famille": "Évaluation", "constat": "Provision insuffisante", "severite": "Élevé"}]})
    revision = _patch_section(api, mission, assistant, "conclusion", {"synthese": "", "decision": "completer"}).json()
    assert revision["statut"] == "4c"
    # Pas de soumission sans « Cycle justifié ».
    assert api.post(_url(mission, "/revision/soumettre"), headers=assistant).status_code == 400


def test_cycle_complet_avec_retour_senior(api, employees, mission):
    assistant = employees["assistant"]["headers"]
    senior = employees["senior"]["headers"]
    manager = employees["manager"]["headers"]
    associe = employees["associe"]["headers"]

    # 1. Soumission -> revue Senior, travaux verrouillés.
    revision = _preparer_et_soumettre(api, mission, assistant)
    assert revision["statut"] == "revue_senior" and revision["verrouille"]
    assert _etats(revision) == {"collaborateur": "fait", "senior": "en_cours", "manager": "en_attente", "associe": "en_attente"}
    assert _patch_section(api, mission, assistant, "lead", {"commentaires": {}}).status_code == 409

    # Le Manager ne peut pas sauter le niveau Senior.
    assert api.post(_url(mission, "/revue/manager/valider"), headers=manager).status_code == 409
    # Un assistant ne peut pas faire de revue.
    assert api.post(_url(mission, "/revue/senior/valider"), headers=assistant).status_code == 403

    # 2. Le Senior émet un point sur la justification et renvoie le cycle.
    point = api.post(
        _url(mission, "/points"),
        headers=senior,
        json={"niveau": "senior", "corps": "Créances > 12 mois non analysées", "cible": {"etape": "4c"}},
    ).json()
    assert point["ref"] == "C-1" and point["statut"] == "ouvert" and point["roundCount"] == 0

    _cocher_checklist(api, mission, "senior", senior)
    blocked = api.post(_url(mission, "/revue/senior/valider"), headers=senior)
    assert blocked.status_code == 409 and "C-1" in blocked.json()["detail"]

    revision = api.post(_url(mission, "/revue/senior/retour"), headers=senior, json={"commentaire": "À compléter"}).json()
    assert revision["statut"] == "4c" and revision["retourDe"] == "senior" and not revision["verrouille"]
    assert _etats(revision)["senior"] == "retour"
    point = api.get(_url(mission, "/points"), headers=assistant).json()[0]
    assert point["roundCount"] == 1 and not point["alerte"]

    # 3. Le collaborateur corrige, répond avec une pièce jointe, et resoumet.
    assert _patch_section(api, mission, assistant, "justification", {"commentaire": "Analyse des créances faite"}).status_code == 200
    response = api.post(
        _url(mission, f"/points/{point['id']}/reponses"),
        headers=assistant,
        data={"texte": "Créances encaissées en janvier N+1", "ptRef": "4d / Balance âgée"},
        files=[("fichiers", ("releve.pdf", b"%PDF-1.4", "application/pdf"))],
    )
    assert response.status_code == 201, response.text
    reponse = response.json()["reponses"][0]
    assert reponse["ptRef"] == "4d / Balance âgée" and reponse["pieces"][0]["fileName"] == "releve.pdf"

    revision = _preparer_et_soumettre(api, mission, assistant)
    assert revision["statut"] == "revue_senior"

    # 4. Le Senior clôture le point et valide -> Manager.
    assert api.patch(_url(mission, f"/points/{point['id']}"), headers=senior, json={"statut": "cleared"}).json()["statut"] == "cleared"
    revision = api.post(_url(mission, "/revue/senior/valider"), headers=senior).json()
    assert revision["statut"] == "revue_manager"
    assert _etats(revision)["senior"] == "fait"

    # 5. Manager puis Associé valident -> clôture du cycle.
    _cocher_checklist(api, mission, "manager", manager)
    assert api.post(_url(mission, "/revue/manager/valider"), headers=manager).json()["statut"] == "revue_associe"
    _cocher_checklist(api, mission, "associe", associe)
    revision = api.post(_url(mission, "/revue/associe/valider"), headers=associe).json()
    assert revision["statut"] == "cloture"
    assert set(_etats(revision).values()) == {"fait"}
    assert [h["action"] for h in revision["historique"]] == ["soumission", "retour", "soumission", "validation", "validation", "validation"]


def test_retour_manager_revient_au_manager_et_conserve_la_validation_senior(api, employees, mission):
    assistant = employees["assistant"]["headers"]
    senior = employees["senior"]["headers"]
    manager = employees["manager"]["headers"]

    _preparer_et_soumettre(api, mission, assistant)
    _cocher_checklist(api, mission, "senior", senior)
    api.post(_url(mission, "/revue/senior/valider"), headers=senior)

    api.post(_url(mission, "/points"), headers=manager, json={"niveau": "manager", "corps": "Variation non expliquée", "cible": {"etape": "4b"}})
    revision = api.post(_url(mission, "/revue/manager/retour"), headers=manager, json={}).json()
    assert revision["statut"] == "4b"

    revision = _preparer_et_soumettre(api, mission, assistant)
    assert revision["statut"] == "revue_manager"
    senior_step = next(e for e in revision["chaine"] if e["niveau"] == "senior")
    assert senior_step["etat"] == "fait" and senior_step["modifieApresValidation"]


def test_alerte_au_dela_de_trois_retours(api, employees, mission):
    assistant = employees["assistant"]["headers"]
    senior = employees["senior"]["headers"]

    _preparer_et_soumettre(api, mission, assistant)
    point = api.post(_url(mission, "/points"), headers=senior, json={"niveau": "senior", "corps": "Pièce manquante", "cible": {"etape": "4c"}}).json()

    for round_attendu in range(1, 5):
        api.post(_url(mission, "/revue/senior/retour"), headers=senior, json={})
        point = api.get(_url(mission, "/points"), headers=senior).json()[0]
        assert point["roundCount"] == round_attendu
        assert point["alerte"] == (round_attendu > 3)
        _preparer_et_soumettre(api, mission, assistant)

    revisions = api.get(f"/missions/{mission}/revisions", headers=senior).json()
    cycle_c = next(r for r in revisions if r["code"] == "C")
    assert cycle_c["alertes"] == 1 and cycle_c["pointsOuverts"] == 1


def test_papier_de_travail_version_journal_et_export(api, employees, mission):
    assistant = employees["assistant"]["headers"]
    base = _url(mission, "/papiers/4d")

    papier = api.get(base, headers=assistant).json()
    assert [f["nom"] for f in papier["feuilles"]] == ["Balance âgée", "Cut-off", "Circularisation"]
    assert papier["mode"] == "integre" and papier["version"] == 0

    feuilles = papier["feuilles"]
    feuilles[0]["data"][0][1] = "168 400 000"
    feuilles[0]["data"][1][1] = "48600000"
    papier = api.patch(base, headers=assistant, json={"version": 0, "feuilles": feuilles}).json()
    assert papier["version"] == 1 and papier["updatedBy"]["nom"] == "Assistant ASSISTANT" and papier["updatedAt"]

    assert api.patch(base, headers=assistant, json={"version": 0, "mode": "lien"}).status_code == 409

    journal = api.get(f"{base}/versions", headers=assistant).json()
    assert [v["version"] for v in journal] == [1, 0]

    export = api.post(f"{base}/export", headers=assistant).json()
    wb = load_workbook(upload_file_path(export["fileUrl"]))
    ws = wb["Balance âgée"]
    assert ws["A1"].value == "Tranche"
    assert ws["B2"].value == 168400000
    # La formule de l'éditeur (ligne 5 = Total) est décalée sous la ligne d'en-têtes.
    assert ws["B6"].value == "=SUM(B2:B5)"


def test_papier_lien_excel_relit_le_fichier(api, employees, mission):
    from openpyxl import Workbook

    assistant = employees["assistant"]["headers"]
    base = _url(mission, "/papiers/4b")
    papier = api.get(base, headers=assistant).json()

    wb = Workbook()
    wb.active.title = "Lead"
    wb.active.append(["Compte", "Solde N"])
    wb.active.append(["411000", 241900000])
    buffer = io.BytesIO()
    wb.save(buffer)

    response = api.post(
        f"{base}/fichier",
        headers=assistant,
        data={"version": str(papier["version"])},
        files={"fichier": ("Lead.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    papier = response.json()
    assert papier["mode"] == "lien"
    assert papier["lien"]["feuilles"][0]["colonnes"] == ["Compte", "Solde N"]
    assert papier["lien"]["feuilles"][0]["data"] == [["411000", 241900000]]

    # Une fois soumis, le PT est en lecture seule.
    _preparer_et_soumettre(api, mission, assistant)
    papier = api.get(base, headers=assistant).json()
    assert papier["verrouille"]
    assert api.patch(base, headers=assistant, json={"version": papier["version"], "titre": "x"}).status_code == 409


def test_droits_renvoyes_selon_l_utilisateur(api, employees, mission):
    revision = api.get(_url(mission, "/revision"), headers=employees["assistant"]["headers"]).json()
    assert revision["droits"] == {"editer": True, "repondre": True, "revue": {"senior": False, "manager": False, "associe": False}}
    revision = api.get(_url(mission, "/revision"), headers=employees["manager"]["headers"]).json()
    assert revision["droits"]["revue"] == {"senior": True, "manager": True, "associe": False}


def test_cycle_seulement_consulte_reste_non_commence(api, employees, mission):
    headers = employees["assistant"]["headers"]
    api.get(f"/missions/{mission}/cycles/A/revision", headers=headers)
    _patch_section(api, mission, headers, "lead", {"commentaires": {"411": "ok"}})
    statuts = {c["code"]: c["statut"] for c in api.get(f"/missions/{mission}/revisions", headers=headers).json()}
    assert statuts["A"] == "non_commence" and statuts["C"] == "4b"
